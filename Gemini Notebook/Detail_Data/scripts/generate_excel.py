import sys
import os
import json
import datetime
from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.worksheet.pagebreak import Break

def get_schema():
    schema_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), "config", "schema.json")
    with open(schema_path, "r", encoding="utf-8") as f:
        return json.load(f)

def generate_excel(input_json_path, output_excel_path, logo_path=None):
    schema = get_schema()
    headers = schema["headers"]
    
    with open(input_json_path, "r", encoding="utf-8") as f:
        data = json.load(f)
        
    wb = Workbook()
    ws = wb.active
    ws.title = "Slip Data"
    
    # Page Setup for A4 Landscape
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = False
    
    # Convert Pixels to Inches (assuming 96 DPI standard)
    # Top 157px, Bottom 179px, Left/Right 102px
    ws.page_margins.top = 157 / 96.0
    ws.page_margins.bottom = 179 / 96.0
    ws.page_margins.left = 102 / 96.0
    ws.page_margins.right = 102 / 96.0
    
    # Set Column Widths
    col_widths = schema.get("column_widths", {})
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    # Styles
    font_bold = Font(name='Tahoma', size=11, bold=True)
    font_normal = Font(name='Tahoma', size=10)
    font_small = Font(name='Tahoma', size=8, color="555555")
    
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='center')
    
    thin = Side(border_style="thin", color="000000")
    border_all = Border(top=thin, left=thin, right=thin, bottom=thin)

    disclaimer = (
        '"DIGITAL EVIDENCE เป็นเพียงการเครื่องมืออำนวยความสะดวกให้กับผู้ว่าจ้าง โดยไม่ได้ดัดแปลง แก้ไข เพิ่ม-ลบ เนื้อหา\n'
        'จากต้นฉบับใดๆ และไม่มีส่วนเกี่ยวข้องใดๆกับเนื้อหาในเอกสาร เป็นเพียงเครื่องมือที่ทำงานเกี่ยวกับระบบไฟล์\n'
        'เอกสารแบบอิเล็กทรอนิกส์ เท่านั้น"'
    )
    timestamp = datetime.datetime.now().strftime("%d/%m/%Y : %H.%M")

    # Chunk data into pages of 20
    chunk_size = 20
    chunks = [data[i:i + chunk_size] for i in range(0, max(len(data), 1), chunk_size)]
    total_pages = len(chunks)
    
    current_row = 1
    
    for page_idx, chunk in enumerate(chunks):
        page_num = page_idx + 1
        
        # --- HEADER BLOCK ---
        # Logo path candidates
        actual_logo = logo_path
        if not actual_logo or not os.path.exists(actual_logo):
            candidates = [
                r"C:\Users\EVE\OneDrive\เดสก์ท็อป\EVIDENCE.png",
                r"C:\Users\EVE\OneDrive\เดสก์ท็อป\EVIDENCE.jpg",
                r"C:\Users\EVE\OneDrive\เดสก์ท็อป\unnamed.png",
                r"C:\Users\EVE\.gemini\config\skills\Dicut_Chat\assets\EVIDENCE.png",
                os.path.join(os.path.dirname(os.path.dirname(__file__)), "config", "logo.png")
            ]
            for c in candidates:
                if os.path.exists(c):
                    actual_logo = c
                    break
            
        if actual_logo and os.path.exists(actual_logo):
            try:
                img = OpenpyxlImage(actual_logo)
                img.width = 100
                img.height = 100
                ws.add_image(img, f"A{current_row}")
            except Exception as e:
                print(f"Warning: Failed to add logo image: {e}")
        
        ws[f"C{current_row}"] = "EVIDENCE"
        ws[f"C{current_row}"].font = font_bold
        
        ws[f"C{current_row+1}"] = f"PAGE: {page_num}"
        ws[f"C{current_row+1}"].font = font_bold
        
        ws[f"C{current_row+2}"] = timestamp
        ws[f"C{current_row+2}"].font = font_bold
        
        ws[f"J{current_row}"] = f"{page_num} / {total_pages}"
        ws[f"J{current_row}"].font = font_bold
        ws[f"J{current_row}"].alignment = align_right
        
        # --- TABLE HEADERS ---
        header_row = current_row + 5
        
        # Table block is 650px height for 21 rows (1 header + 20 data) = 30.95px per row.
        # Convert pixel to points (1px = 0.75 points) => 30.95 * 0.75 = 23.2 points
        table_row_height_pts = 23.2
        ws.row_dimensions[header_row].height = table_row_height_pts
        
        for col_idx, header_text in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col_idx, value=header_text)
            cell.font = font_bold
            cell.alignment = align_center
            cell.border = border_all
            
        # --- TABLE DATA ---
        data_start_row = header_row + 1
        for row_offset in range(20):
            data_row = data_start_row + row_offset
            ws.row_dimensions[data_row].height = table_row_height_pts
            # If we have data for this row
            if row_offset < len(chunk):
                item = chunk[row_offset]
                if isinstance(item, dict):
                    row_data = [
                        row_offset + 1 + (page_idx * chunk_size), # ลำดับ
                        item.get("date", ""),
                        item.get("time", ""),
                        item.get("sender_bank", ""),
                        item.get("sender_name", ""),
                        item.get("amount", ""),
                        item.get("receiver_name", ""),
                        item.get("receiver_bank", ""),
                        item.get("memo", ""),
                        item.get("remarks", "")
                    ]
                else:
                    row_data = item
            else:
                # Empty row to fill the 20-row grid
                row_data = [""] * len(headers)
                
            for col_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=data_row, column=col_idx, value=val)
                cell.font = font_normal
                cell.alignment = align_center
                cell.border = border_all
                
        # --- FOOTER BLOCK ---
        footer_row = data_start_row + 21
        ws.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=10)
        footer_cell = ws.cell(row=footer_row, column=1, value=disclaimer)
        footer_cell.font = font_small
        footer_cell.alignment = align_center
        
        # --- PAGE BREAK ---
        # Add a horizontal page break after this chunk so it prints on a new page
        page_break_row = footer_row + 1
        ws.row_breaks.append(Break(id=page_break_row))
        
        # Move to next block start
        current_row = page_break_row + 2

    wb.save(output_excel_path)
    try:
        print(f"Successfully generated {os.path.basename(output_excel_path)} with {total_pages} pages.")
    except Exception:
        pass

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python generate_excel.py <input_json> <output_excel> [logo_image_path]")
        sys.exit(1)
        
    input_json = sys.argv[1]
    output_excel = sys.argv[2]
    logo = sys.argv[3] if len(sys.argv) > 3 else None
    
    generate_excel(input_json, output_excel, logo)
