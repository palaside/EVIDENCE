#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
tools/generate_slip_evidence_dossier.py
======================================
Generates the Official Court-Admissible Forensic Evidence Dossier:
1. 95 Individual Slip Pages (A4 Portrait, 993 x 1406 px)
   - Strict adherence to `_skills/slip-block-fit/SKILL.md` (645x890 block canvas centered at (322.5, 445))
   - Header Ribbon (EVIDENCE Logo, MODE: SLIP, CORROBORATED Details, PAGE Stamp)
   - Footer 3-line Forensic Disclaimer
2. 5 Landscape Summary Statement Pages (A4 Landscape, 1406 x 993 px)
   - 10-Column Sarabun Table (1,202 px wide, centered text in all cells)
   - 20 rows per page
Total: Exactly 100 pages of complete forensic evidence.
"""

import os
import sys
import json
import datetime

if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if sys.stderr and hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

from PIL import Image, ImageDraw, ImageFont, ImageChops

import pypdfium2 as pdfium
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
OUTPUT_DIR = os.path.join(PROJECT_ROOT, "Folder_Out")

# Candidate resolution for MASTER_PDF and SLIP_INDEX_JSON
import glob
def _resolve_file(pattern, default_path):
    if os.path.exists(default_path):
        return default_path
    matches = glob.glob(os.path.join(OUTPUT_DIR, "**", pattern), recursive=True)
    if matches:
        return matches[0]
    return default_path

MASTER_PDF = _resolve_file("*Evidence_Chat_Master_Combined_Vol1_to_3.pdf", os.path.join(OUTPUT_DIR, "Evidence_Chat_Master_Combined_Vol1_to_3.pdf"))
SLIP_INDEX_JSON = _resolve_file("*Evidence_Chat_Master_Combined_Vol1_to_3_Slip_Index.json", os.path.join(OUTPUT_DIR, "Evidence_Chat_Master_Combined_Vol1_to_3_Slip_Index.json"))
OUTPUT_PDF = os.path.join(OUTPUT_DIR, "Evidence_Slips_95_Fit_Summary.pdf")
OUTPUT_XLSX = os.path.join(OUTPUT_DIR, "Evidence_Slips_95_Statement.xlsx")
LOGO_PATH = os.path.join(PROJECT_ROOT, "EVIDENCE.png")

# --- Import Central SSOT Evidence Theme ---
try:
    from core.evidence_theme import (
        load_evidence_fonts,
        load_evidence_logo,
        apply_header_ribbon,
        apply_footer_disclaimer,
        extract_pure_slip_card,
        fit_slip_block,
        DISCLAIMER_LINES,
    )
except ImportError:
    try:
        from evidence_theme import (
            load_evidence_fonts,
            load_evidence_logo,
            apply_header_ribbon,
            apply_footer_disclaimer,
            extract_pure_slip_card,
            fit_slip_block,
            DISCLAIMER_LINES,
        )
    except ImportError:
        sys.path.insert(0, os.path.dirname(__file__))
        sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "core"))
        from evidence_theme import (
            load_evidence_fonts,
            load_evidence_logo,
            apply_header_ribbon,
            apply_footer_disclaimer,
            extract_pure_slip_card,
            fit_slip_block,
            DISCLAIMER_LINES,
        )

# Target Evidence Block standard
TW, TH = 645, 890


def render_portrait_slip_page(slip_block, item_info, page_num, total_pages, fonts, logo_img):
    pw, ph = 993, 1406
    canvas = Image.new("RGB", (pw, ph), "white")

    # Center 645x890 block onto A4 Portrait
    paste_x = (pw - TW) // 2   # 174
    paste_y = 120 + ((1200 - TH) // 2)  # 275
    canvas.paste(slip_block, (paste_x, paste_y))

    # 1. Header Evidence Ribbon via SSOT Theme
    idx = item_info.get("index", page_num)
    amt = item_info.get("amount", "0.00")
    dt = item_info.get("datetime") or item_info.get("date_time", "-")
    corrob_str = f"สลิปหลักฐานโอนเงิน ลำดับที่ {idx:02d} | ยอดเงิน: {amt} บาท ({dt})"
    page_str = f"{page_num} / {total_pages}"

    apply_header_ribbon(
        canvas=canvas,
        mode="SLIP",
        corroborated=corrob_str,
        page_str=page_str,
        fonts=fonts,
        logo_img=logo_img,
        is_landscape=False
    )

    # 2. Footer Legal Disclaimer via SSOT Theme
    apply_footer_disclaimer(
        canvas=canvas,
        fonts=fonts,
        is_landscape=False
    )

    return canvas


def render_landscape_summary_page(rows_chunk, start_idx, page_num, total_pages, fonts, logo_img):
    lw, lh = 1406, 993
    canvas = Image.new("RGB", (lw, lh), "white")
    draw = ImageDraw.Draw(canvas)

    # 1. Header Ribbon via SSOT Theme
    corrob_str = f"ตารางสรุปรายการธุรกรรมทางการเงิน (สลิปหลักฐาน ลำดับที่ {start_idx} - {start_idx + len(rows_chunk) - 1})"
    page_str = f"{page_num} / {total_pages}"

    apply_header_ribbon(
        canvas=canvas,
        mode="STATEMENT / 10-COLUMN SUMMARY TABLE",
        corroborated=corrob_str,
        page_str=page_str,
        fonts=fonts,
        logo_img=logo_img,
        is_landscape=True
    )

    # 2. Table Render (1,202 px wide, margins 102 px)
    start_x = 102
    start_y = 115
    col_w = [52, 105, 75, 125, 180, 110, 180, 125, 125, 125]  # sum = 1202
    headers = ["ลำดับ", "วันที่", "เวลา", "ธนาคารผู้โอน", "ชื่อผู้โอน", "จำนวนเงิน", "ชื่อผู้รับ", "ธนาคารผู้รับ", "บันทึกช่วยจำ", "หมายเหตุ (Ref)"]
    row_h = 32

    # Draw Header Row
    cur_x = start_x
    for i, (h_title, cw) in enumerate(zip(headers, col_w)):
        draw.rectangle([cur_x, start_y, cur_x + cw, start_y + row_h], fill="#1E3A8A", outline="#CBD5E1")
        bbox = draw.textbbox((0, 0), h_title, font=fonts["table_hdr"])
        tw = bbox[2] - bbox[0]
        th = bbox[3] - bbox[1]
        draw.text((cur_x + (cw - tw) // 2, start_y + (row_h - th) // 2), h_title, fill="#FFFFFF", font=fonts["table_hdr"])
        cur_x += cw

    # Draw Data Rows (Center-Aligned)
    for r_idx, item in enumerate(rows_chunk):
        y_top = start_y + ((r_idx + 1) * row_h)
        cur_x = start_x
        row_bg = "#F8FAFC" if (r_idx % 2 == 1) else "#FFFFFF"

        # Split date and time
        dt_full = item.get("datetime") or item.get("date_time", "-")
        if " - " in dt_full:
            d_part, t_part = dt_full.split(" - ", 1)
        else:
            d_part, t_part = dt_full, "-"

        row_vals = [
            f"{item.get('index', start_idx + r_idx):02d}",
            d_part.strip(),
            t_part.strip(),
            item.get("sender_bank", "-"),
            item.get("sender_name", "-"),
            f"{item.get('amount', '0.00')} บาท",
            item.get("receiver_name", "-"),
            item.get("receiver_bank", "-"),
            item.get("memo", "-") or "-",
            item.get("ref_id", "-")
        ]

        for i, (val, cw) in enumerate(zip(row_vals, col_w)):
            draw.rectangle([cur_x, y_top, cur_x + cw, y_top + row_h], fill=row_bg, outline="#E2E8F0")
            font_use = fonts["table_small"] if (len(val) > 22 or i == 9) else fonts["table_body"]
            
            # Shorten display for long ref ID
            display_val = val if len(val) <= 18 else (val[:8] + ".." + val[-7:]) if i == 9 else val
            bbox = draw.textbbox((0, 0), display_val, font=font_use)
            tw = bbox[2] - bbox[0]
            th = bbox[3] - bbox[1]
            tx = max(cur_x + 2, cur_x + (cw - tw) // 2)
            ty = y_top + (row_h - th) // 2
            draw.text((tx, ty), display_val, fill="#0F172A", font=font_use)
            cur_x += cw

    # 3. Footer Legal Disclaimer via SSOT Theme
    apply_footer_disclaimer(
        canvas=canvas,
        fonts=fonts,
        is_landscape=True
    )

    return canvas

def export_evidence_excel(slips, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "สรุปสลิปหลักฐาน_95_รายการ"

    headers = [
        "ลำดับ", "หน้า PDF ต้นฉบับ", "วันที่", "เวลา",
        "ธนาคารผู้โอน", "ชื่อผู้โอน", "จำนวนเงิน (บาท)",
        "ชื่อผู้รับเงิน", "ธนาคารผู้รับ", "บันทึกช่วยจำ", "Transaction Ref (QR)"
    ]
    ws.append(headers)

    f_hdr = Font(name="Sarabun", size=11, bold=True, color="FFFFFF")
    f_body = Font(name="Sarabun", size=10)
    fill_hdr = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.font = f_hdr
        c.fill = fill_hdr
        c.alignment = align_center

    for idx, s in enumerate(slips):
        dt_full = s.get("datetime") or s.get("date_time", "-")
        if " - " in dt_full:
            d_part, t_part = dt_full.split(" - ", 1)
        else:
            d_part, t_part = dt_full, "-"

        amt_str = s.get("amount", "0").replace(",", "").strip()
        try:
            amt_num = float(amt_str)
        except Exception:
            amt_num = 0.0

        row = [
            s.get("index", idx + 1),
            f"p.{s.get('page', '-')}",
            d_part.strip(),
            t_part.strip(),
            s.get("sender_bank", "-"),
            s.get("sender_name", "-"),
            amt_num,
            s.get("receiver_name", "-"),
            s.get("receiver_bank", "-"),
            s.get("memo", "-") or "-",
            s.get("ref_id", "-")
        ]
        ws.append(row)
        for col in range(1, len(row) + 1):
            c = ws.cell(row=idx + 2, column=col)
            c.font = f_body
            c.alignment = align_center
            c.border = thin_border
            if col == 7:
                c.number_format = '#,##0.00'

    for col in ws.columns:
        max_len = max(len(str(c.value or '')) for c in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)

    wb.save(out_path)
    print(f"✅ Excel Exported: {out_path}")

def main():
    import argparse
    parser = argparse.ArgumentParser(description="Generate Official Court Slip Evidence Dossier")
    parser.add_argument("--sample", action="store_true", help="Generate exactly 1 real page sample preview")
    parser.add_argument("--all", action="store_true", help="Generate full complete dossier (all pages)")
    args = parser.parse_args()

    print("=" * 80)
    print("🏛️ FORENSIC PIPELINE: Generating Official Court Slip Evidence Dossier")
    print("   Standard: _skills/slip-block-fit (645x890 block canvas centered at (322.5, 445))")
    print("   Theme   : Central SSOT evidence_theme (Evidence Ribbon + 3-line Disclaimer)")
    print("=" * 80)

    if not os.path.exists(SLIP_INDEX_JSON):
        print(f"❌ Error: Slip index not found at {SLIP_INDEX_JSON}")
        sys.exit(1)
    if not os.path.exists(MASTER_PDF):
        print(f"❌ Error: Master PDF not found at {MASTER_PDF}")
        sys.exit(1)

    is_sample = args.sample
    if not args.sample and not args.all:
        if sys.stdin and sys.stdin.isatty():
            choice = input("\nต้องการสร้าง 'ตัวอย่าง (1 หน้าจริง)' หรือ 'ทั้งหมด' (sample/all) [default: sample]: ").strip().lower()
            if choice in ("all", "ท", "ทั้งหมด"):
                is_sample = False
            else:
                is_sample = True
        else:
            is_sample = False

    with open(SLIP_INDEX_JSON, "r", encoding="utf-8") as f:
        slips = json.load(f)

    if is_sample:
        print("\n🔍 MODE: [SAMPLE PREVIEW] Exactly 1 real page will be generated.")
        slips_to_process = slips[:1]
        summary_chunks = []
        total_dossier_pages = 1
        output_pdf_path = os.path.join(OUTPUT_DIR, "SAMPLE_Evidence_Slips_95_Fit_Summary.pdf")
    else:
        print(f"[*] Total Slips to Process: {len(slips)}")
        # 1. Export Excel Statement
        export_evidence_excel(slips, OUTPUT_XLSX)
        slips_to_process = slips
        rows_per_page = 20
        summary_chunks = [slips[i:i + rows_per_page] for i in range(0, len(slips), rows_per_page)]
        total_dossier_pages = len(slips) + len(summary_chunks)
        output_pdf_path = OUTPUT_PDF

    # 2. Setup Resources via SSOT Theme
    fonts = load_evidence_fonts()
    logo_img = load_evidence_logo(target_h=75)
    if logo_img:
        print("[*] Loaded Evidence Logo (height 75px)")

    print(f"[*] Total Dossier Pages to Generate: {total_dossier_pages}")

    pdf_doc = pdfium.PdfDocument(MASTER_PDF)
    rendered_pages = []

    # 3. Generate Slip Pages via pure slip card isolation + slip-block-fit
    print("\n--- Part 1: Generating Slip Pages (Pure Slip Isolation + slip-block-fit 645x890) ---")
    for idx, item in enumerate(slips_to_process):
        p_num = item["page"]
        page_obj = pdf_doc[p_num - 1]
        
        # Render at 1.5x resolution for crystal clear pixel quality
        bmp = page_obj.render(scale=1.5)
        page_pil = bmp.to_pil()
        pw, ph = page_pil.size

        # Extract central slip card ROI (between y=10% and y=90%)
        roi = page_pil.crop((0, int(ph * 0.10), pw, int(ph * 0.90)))

        # Pure Slip Card Isolation: Cut out surrounding chat messages & conversation context 100%
        slip_card = extract_pure_slip_card(roi)

        # Apply slip-block-fit protocol (645x890, centered at (322.5, 445))
        block_canvas = fit_slip_block(slip_card)

        # Render complete Portrait A4 page
        dossier_page_idx = idx + 1
        page_canvas = render_portrait_slip_page(
            block_canvas, item, dossier_page_idx, total_dossier_pages, fonts, logo_img
        )
        rendered_pages.append(page_canvas)

        if (idx + 1) % 15 == 0 or (idx + 1) == len(slips_to_process):
            print(f"    [{idx + 1:02d}/{len(slips_to_process)}] Slip pages generated and block-fitted.")

    pdf_doc.close()

    # 4. Generate Landscape Summary Statement Pages (if not sample)
    if summary_chunks:
        print("\n--- Part 2: Generating Landscape Summary Statement Pages ---")
        for c_idx, chunk in enumerate(summary_chunks):
            start_row_num = (c_idx * rows_per_page) + 1
            dossier_page_idx = len(slips) + c_idx + 1
            land_canvas = render_landscape_summary_page(
                chunk, start_row_num, dossier_page_idx, total_dossier_pages, fonts, logo_img
            )
            rendered_pages.append(land_canvas)
            print(f"    Table Page {c_idx + 1} (Rows {start_row_num} - {start_row_num + len(chunk) - 1}) generated.")

    # 5. Save Dossier PDF & Sample Preview Image
    print(f"\n[*] Saving Dossier PDF: {output_pdf_path}...")
    first_page = rendered_pages[0]
    rest_pages = rendered_pages[1:]
    first_page.save(
        output_pdf_path,
        "PDF",
        resolution=150.0,
        save_all=True,
        append_images=rest_pages
    )

    if is_sample:
        sample_png_path = output_pdf_path.replace(".pdf", ".png")
        first_page.save(sample_png_path, "PNG")
        print(f"   [OK] Generated 1-Page Real Sample Preview Image: {sample_png_path}")

    print("=" * 80)
    print(f"🎉 SUCCESS! Evidence Dossier Generated: {output_pdf_path}")
    print(f"   - File Size: {os.path.getsize(output_pdf_path) / (1024*1024):.2f} MB")
    print(f"   - Total Pages: {len(rendered_pages)} pages")
    print("=" * 80)

if __name__ == "__main__":
    main()
